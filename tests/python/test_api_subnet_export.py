import io
from openpyxl import load_workbook


def _cell_by_value(ws, value):
    return next(c for row in ws.iter_rows() for c in row if c.value == value)


def test_export_xlsx_colores_y_header(client, seed):
    r = client.get(f"/api/v1/subnets/{seed['subnet'].id}/export")
    assert r.status_code == 200
    ws = load_workbook(io.BytesIO(r.content)).active
    assert ws["A1"].value == "192.168.1.0/24 Red de prueba"
    assert _cell_by_value(ws, 10).fill.fgColor.rgb.endswith("198754")  # FREE
    assert _cell_by_value(ws, 11).fill.fgColor.rgb.endswith("DC3545")  # ASSIGNED


def test_export_404_y_subred_grande_400(client):
    assert client.get("/api/v1/subnets/9999/export").status_code == 404
    sid = client.post("/api/v1/subnets", json={"cidr": "10.0.0.0/8", "name": "Grande"}).json()["id"]
    assert client.get(f"/api/v1/subnets/{sid}/export").status_code == 400