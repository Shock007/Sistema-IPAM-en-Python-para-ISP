"""Reporte XLSX de una subred: header = subred, celdas = hosts coloreados por estado."""
import io
import ipaddress
from math import ceil

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from app.Models.ip_address import IPStatus

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
ROWS_PER_COLUMN = 52
MAX_HOSTS = 4096

# (fondo, texto): mismos colores del dashboard/SPA.
COLORS = {
    IPStatus.FREE: ("198754", "FFFFFF"),
    IPStatus.ASSIGNED: ("DC3545", "FFFFFF"),
    IPStatus.ACTIVE: ("FFC107", "212529"),
}
LABELS = {
    IPStatus.FREE: "Libre",
    IPStatus.ASSIGNED: "Asignada",
    IPStatus.ACTIVE: "Activa sin registrar",
}


class SubnetReportError(ValueError):
    pass


def _host_label(ip: ipaddress.IPv4Address, prefix: int):
    k = 1 if prefix >= 24 else 2 if prefix >= 16 else 3
    parts = str(ip).split(".")[-k:]
    return int(parts[0]) if k == 1 else ".".join(parts)


def build_subnet_report(subnet, ips) -> bytes:
    try:
        net = ipaddress.ip_network(subnet.cidr, strict=False)
    except ValueError as exc:
        raise SubnetReportError(f"CIDR inválido: {subnet.cidr}") from exc
    if net.version != 4:
        raise SubnetReportError("La exportación solo soporta subredes IPv4.")
    if net.num_addresses > MAX_HOSTS + 2:
        raise SubnetReportError(
            f"La subred tiene {net.num_addresses} direcciones; máximo exportable: {MAX_HOSTS}."
        )

    hosts = list(net.hosts())
    status_by_ip = {i.ip_address: i.status for i in ips}
    ncols = max(1, ceil(len(hosts) / ROWS_PER_COLUMN))

    wb = Workbook()
    ws = wb.active
    ws.title = "Reporte IP"
    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center")

    # Header: subred completa
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    head = ws.cell(row=1, column=1, value=f"{subnet.cidr} {subnet.name}")
    head.font = Font(bold=True, color="FFFFFF", size=12)
    head.fill = PatternFill("solid", fgColor="1F4E78")
    head.alignment = center
    ws.row_dimensions[1].height = 22

    for c in range(1, ncols + 1):
        ws.column_dimensions[ws.cell(row=1, column=c).column_letter].width = 12

    # Hosts: llenado por columnas (52 filas c/u)
    for idx, host in enumerate(hosts):
        cell = ws.cell(row=2 + idx % ROWS_PER_COLUMN, column=1 + idx // ROWS_PER_COLUMN,
                       value=_host_label(host, net.prefixlen))
        cell.alignment = center
        cell.border = border
        status = status_by_ip.get(str(host))
        if status in COLORS:
            bg, fg = COLORS[status]
            cell.fill = PatternFill("solid", fgColor=bg)
            cell.font = Font(color=fg)

    # Leyenda a la derecha
    sw, tx = ncols + 2, ncols + 3
    ws.column_dimensions[ws.cell(row=1, column=tx).column_letter].width = 26
    for n, st in enumerate(IPStatus):
        bg, _ = COLORS[st]
        ws.cell(row=2 + n, column=sw).fill = PatternFill("solid", fgColor=bg)
        ws.cell(row=2 + n, column=tx, value=LABELS[st])
    ws.cell(row=2 + len(IPStatus), column=tx, value="Sin color: no registrada en el sistema")

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()